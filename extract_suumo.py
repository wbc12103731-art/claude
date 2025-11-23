#!/usr/bin/env python3
"""
提供されたSUUMO HTMLから物件情報を抽出してExcelに出力
"""

from bs4 import BeautifulSoup
import pandas as pd
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

def extract_properties_from_html(html_file):
    """HTMLファイルから物件情報を抽出"""

    with open(html_file, 'r', encoding='utf-8') as f:
        html = f.read()

    soup = BeautifulSoup(html, 'html.parser')

    properties = []

    # 物件ユニットを取得
    property_units = soup.select('.property_unit')
    print(f"物件ユニット数: {len(property_units)}")

    for unit in property_units:
        prop = {}

        # 物件名
        title_elem = unit.select_one('.property_unit-title_wide a')
        if title_elem:
            prop['物件名'] = title_elem.get_text(strip=True)
            prop['リンク'] = 'https://suumo.jp' + title_elem.get('href', '')

        # dottableからデータを抽出
        dottable_lines = unit.select('.dottable-line table tr')

        for row in dottable_lines:
            cells = row.select('td')
            for cell in cells:
                dl = cell.select_one('dl')
                if dl:
                    dt = dl.select_one('dt')
                    dd = dl.select_one('dd')
                    if dt and dd:
                        label = dt.get_text(strip=True)
                        value = dd.get_text(strip=True)

                        # supタグの処理（m2 → m²）
                        value = value.replace('m2', 'm²')

                        if label and value and label != '\xa0':
                            # ラベルのマッピング
                            label_map = {
                                '販売価格': '価格',
                                '沿線・駅': '交通',
                            }
                            label = label_map.get(label, label)
                            prop[label] = value

        # PRコメントを取得
        lead = unit.select_one('.dottable-lead td')
        if lead:
            prop['PRコメント'] = lead.get_text(strip=True)

        # ピクト（即引渡可、新着など）を取得
        pcts = unit.select('.ui-pct')
        if pcts:
            pct_texts = [pct.get_text(strip=True) for pct in pcts]
            prop['タグ'] = ', '.join(pct_texts)

        if prop.get('物件名'):
            properties.append(prop)

    return properties

def save_to_excel(properties, filename):
    """Excelに保存"""

    if not properties:
        print("保存する物件がありません")
        return False

    # DataFrameを作成
    df = pd.DataFrame(properties)

    # 列の順序を整理
    preferred_order = ['物件名', '価格', '所在地', '交通', '間取り', '専有面積', 'バルコニー', '築年月', 'タグ', 'PRコメント', 'リンク']
    cols = [c for c in preferred_order if c in df.columns]
    other_cols = [c for c in df.columns if c not in preferred_order]
    df = df[cols + other_cols]

    # Excelワークブックを作成
    wb = Workbook()
    ws = wb.active
    ws.title = "物件一覧"

    # スタイル定義
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    cell_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    price_alignment = Alignment(horizontal="right", vertical="center")

    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # データを書き込み
    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.border = thin_border

            if r_idx == 1:  # ヘッダー行
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
            else:  # データ行
                cell.alignment = cell_alignment
                # 価格列は右揃え
                if c_idx == 2:
                    cell.alignment = price_alignment

    # 列幅を調整
    column_widths = {
        'A': 35,  # 物件名
        'B': 12,  # 価格
        'C': 30,  # 所在地
        'D': 35,  # 交通
        'E': 15,  # 間取り
        'F': 18,  # 専有面積
        'G': 12,  # バルコニー
        'H': 12,  # 築年月
        'I': 20,  # タグ
        'J': 50,  # PRコメント
        'K': 60,  # リンク
    }

    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width

    # 行の高さを設定
    ws.row_dimensions[1].height = 25
    for row in range(2, len(properties) + 2):
        ws.row_dimensions[row].height = 30

    # フィルターを設定
    ws.auto_filter.ref = ws.dimensions

    # ウィンドウ枠の固定
    ws.freeze_panes = 'A2'

    # 保存
    wb.save(filename)
    print(f"\nExcelファイルを保存しました: {filename}")
    print(f"物件数: {len(properties)}件")

    return True

def main():
    html_file = '/home/user/claude/suumo_source.html'
    output_file = '/home/user/claude/suumo_properties.xlsx'

    print("物件情報を抽出中...")
    properties = extract_properties_from_html(html_file)

    if properties:
        save_to_excel(properties, output_file)

        # サマリーを表示
        print("\n=== 抽出された物件一覧 ===")
        for i, p in enumerate(properties, 1):
            name = p.get('物件名', '')[:30]
            price = p.get('価格', '')
            area = p.get('専有面積', '')
            layout = p.get('間取り', '')
            print(f"{i:2}. {name} - {price} ({layout}, {area})")
    else:
        print("物件情報を抽出できませんでした")

if __name__ == '__main__':
    main()
