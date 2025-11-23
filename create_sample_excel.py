#!/usr/bin/env python3
"""
SUUMOの物件データサンプルをExcelに出力
実際のデータが手に入ったら parse_suumo_html.py で更新可能
"""

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils.dataframe import dataframe_to_rows

# 大阪市内の中古マンションのサンプルデータ（検索条件に基づく想定）
# 検索条件: 大阪市北区・中央区・西区・天王寺区・福島区・浪速区
# 価格: 1万円〜7500万円, 面積: 50m²以上, 駅徒歩5分以内, 2LDK・3LDK

sample_properties = [
    {
        "物件名": "プラウド大阪城公園",
        "価格": "6,980万円",
        "所在地": "大阪府大阪市中央区森ノ宮中央1丁目",
        "交通": "大阪メトロ中央線「森ノ宮」駅 徒歩3分",
        "間取り": "3LDK",
        "専有面積": "72.5m²",
        "バルコニー": "12.3m²",
        "築年数": "築8年",
        "階数": "15階/20階建",
        "管理費": "15,800円/月",
        "修繕積立金": "12,500円/月",
        "向き": "南",
        "構造": "RC造",
    },
    {
        "物件名": "グランドメゾン梅田タワー",
        "価格": "7,480万円",
        "所在地": "大阪府大阪市北区大淀南2丁目",
        "交通": "JR大阪環状線「福島」駅 徒歩5分",
        "間取り": "3LDK",
        "専有面積": "85.2m²",
        "バルコニー": "15.6m²",
        "築年数": "築5年",
        "階数": "22階/35階建",
        "管理費": "22,000円/月",
        "修繕積立金": "18,000円/月",
        "向き": "南西",
        "構造": "RC造",
    },
    {
        "物件名": "シティタワー大阪本町",
        "価格": "5,980万円",
        "所在地": "大阪府大阪市中央区本町4丁目",
        "交通": "大阪メトロ御堂筋線「本町」駅 徒歩2分",
        "間取り": "2LDK",
        "専有面積": "65.8m²",
        "バルコニー": "10.2m²",
        "築年数": "築12年",
        "階数": "18階/28階建",
        "管理費": "18,500円/月",
        "修繕積立金": "14,200円/月",
        "向き": "東",
        "構造": "SRC造",
    },
    {
        "物件名": "ブランズタワー天王寺",
        "価格": "6,280万円",
        "所在地": "大阪府大阪市天王寺区悲田院町",
        "交通": "JR大阪環状線「天王寺」駅 徒歩4分",
        "間取り": "3LDK",
        "専有面積": "78.3m²",
        "バルコニー": "13.8m²",
        "築年数": "築10年",
        "階数": "25階/32階建",
        "管理費": "20,000円/月",
        "修繕積立金": "16,500円/月",
        "向き": "南",
        "構造": "RC造",
    },
    {
        "物件名": "ライオンズ堀江レジデンス",
        "価格": "4,980万円",
        "所在地": "大阪府大阪市西区北堀江1丁目",
        "交通": "大阪メトロ四つ橋線「四ツ橋」駅 徒歩3分",
        "間取り": "2LDK",
        "専有面積": "58.6m²",
        "バルコニー": "8.5m²",
        "築年数": "築15年",
        "階数": "8階/14階建",
        "管理費": "12,000円/月",
        "修繕積立金": "10,800円/月",
        "向き": "西",
        "構造": "RC造",
    },
    {
        "物件名": "パークタワー西梅田",
        "価格": "7,200万円",
        "所在地": "大阪府大阪市福島区福島5丁目",
        "交通": "JR東西線「新福島」駅 徒歩2分",
        "間取り": "3LDK",
        "専有面積": "82.1m²",
        "バルコニー": "14.2m²",
        "築年数": "築7年",
        "階数": "28階/38階建",
        "管理費": "25,000円/月",
        "修繕積立金": "20,000円/月",
        "向き": "南東",
        "構造": "RC造",
    },
    {
        "物件名": "ジオ難波プレミア",
        "価格": "5,480万円",
        "所在地": "大阪府大阪市浪速区難波中2丁目",
        "交通": "大阪メトロ御堂筋線「なんば」駅 徒歩5分",
        "間取り": "2LDK",
        "専有面積": "62.4m²",
        "バルコニー": "9.8m²",
        "築年数": "築11年",
        "階数": "12階/18階建",
        "管理費": "14,500円/月",
        "修繕積立金": "11,200円/月",
        "向き": "南",
        "構造": "RC造",
    },
    {
        "物件名": "サンクタス北浜タワー",
        "価格": "6,680万円",
        "所在地": "大阪府大阪市中央区北浜2丁目",
        "交通": "大阪メトロ堺筋線「北浜」駅 徒歩1分",
        "間取り": "3LDK",
        "専有面積": "75.9m²",
        "バルコニー": "11.5m²",
        "築年数": "築9年",
        "階数": "20階/25階建",
        "管理費": "19,000円/月",
        "修繕積立金": "15,800円/月",
        "向き": "北",
        "構造": "SRC造",
    },
]

def create_excel(data, filename):
    """物件データをExcelファイルに出力"""

    df = pd.DataFrame(data)

    # Excelワークブックを作成
    wb = Workbook()
    ws = wb.active
    ws.title = "物件一覧"

    # スタイル定義
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    cell_alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    price_alignment = Alignment(horizontal="right", vertical="center")

    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )

    # データフレームをワークシートに書き込み
    for r_idx, row in enumerate(dataframe_to_rows(df, index=False, header=True), 1):
        for c_idx, value in enumerate(row, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=value)
            cell.border = thin_border

            if r_idx == 1:  # ヘッダー行
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment
            else:  # データ行
                cell.alignment = cell_alignment
                # 価格列は右揃え
                if c_idx == 2:
                    cell.alignment = price_alignment

    # 列幅を調整
    column_widths = {
        'A': 30,  # 物件名
        'B': 12,  # 価格
        'C': 30,  # 所在地
        'D': 35,  # 交通
        'E': 10,  # 間取り
        'F': 12,  # 専有面積
        'G': 12,  # バルコニー
        'H': 10,  # 築年数
        'I': 14,  # 階数
        'J': 14,  # 管理費
        'K': 14,  # 修繕積立金
        'L': 8,   # 向き
        'M': 8,   # 構造
    }

    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width

    # 行の高さを設定
    ws.row_dimensions[1].height = 25  # ヘッダー
    for row in range(2, len(data) + 2):
        ws.row_dimensions[row].height = 20

    # フィルターを設定
    ws.auto_filter.ref = ws.dimensions

    # ウィンドウ枠の固定（ヘッダー行を固定）
    ws.freeze_panes = 'A2'

    # 保存
    wb.save(filename)
    print(f"Excelファイルを作成しました: {filename}")
    print(f"物件数: {len(data)}件")

    # 内容のサマリーを表示
    print("\n=== 物件サマリー ===")
    for i, prop in enumerate(data[:5], 1):
        print(f"{i}. {prop['物件名']} - {prop['価格']} ({prop['間取り']}, {prop['専有面積']})")
    if len(data) > 5:
        print(f"   ... 他 {len(data) - 5}件")

if __name__ == '__main__':
    output_file = '/home/user/claude/suumo_properties.xlsx'
    create_excel(sample_properties, output_file)

    print("\n" + "="*60)
    print("注意: これはサンプルデータです。")
    print("実際のSUUMOデータを取得するには:")
    print("1. ブラウザでSUUMOページを開く")
    print("2. Ctrl+U でソースを表示してコピー")
    print("3. テキストファイルとして保存（例: suumo.html）")
    print("4. python parse_suumo_html.py suumo.html")
    print("="*60)
